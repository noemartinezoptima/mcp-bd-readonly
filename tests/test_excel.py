import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(__file__) + "/../src")

from openpyxl import Workbook, load_workbook

from mcp_bd_readonly import tools_excel


def make_workbook(path, with_formulas=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliacion"
    ws["A1"] = "Cliente"
    ws["B1"] = "Base"
    ws["C1"] = "IVA"
    ws["D1"] = "Total"
    ws["A2"] = "MANGO"
    ws["B2"] = 1000
    ws["C2"] = 210
    ws["D2"] = "=B2+C2"
    ws["A3"] = "KIKO"
    ws["B3"] = 2000
    ws["C3"] = 420
    ws["D3"] = "=B3+C3"
    ws["D5"] = "=SUM(D2:D3)"
    ws.merge_cells("F1:G2")
    ws["F1"] = "nota"
    ws2 = wb.create_sheet("Totales")
    ws2["A1"] = "Gran total"
    ws2["B1"] = "=Conciliacion!D5"
    wb.save(path)


class TestLeerExcel(unittest.TestCase):
    def setUp(self):
        import subprocess

        self.tmp = subprocess.run(
            ["mktemp", "-d"], capture_output=True, text=True, check=True
        ).stdout.strip()
        self.xlsx = os.path.join(self.tmp, "muestra.xlsx")
        make_workbook(self.xlsx)
        self.orig_suffix = tools_excel.MAX_FILE_MB

    def tearDown(self):
        import shutil

        tools_excel.MAX_FILE_MB = self.orig_suffix
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_lee_celdas_con_formulas(self):
        r = tools_excel.leer_excel(self.xlsx, include_formulas=True,
                                   force_recalc=False, max_cells=500)
        assert r.get("error") is None, r.get("error")
        assert r["has_formulas"] is True
        assert r["formula_count"] >= 3
        ws = r["sheets"][0]
        assert ws["merged_ranges"] == ["F1:G2"]
        by_ref = {c["ref"]: c for c in ws["cells"]}
        assert by_ref["D2"]["f"] == "=B2+C2"
        assert by_ref["A2"]["v"] == "MANGO"
        assert by_ref["B2"]["v"] == 1000

    def test_original_no_se_modifica(self):
        tools_excel.leer_excel(self.xlsx, include_formulas=True,
                               force_recalc=True, max_cells=500)
        after = load_workbook(self.xlsx, data_only=False)["Conciliacion"]
        assert after["D5"].value == "=SUM(D2:D3)"

    def test_sheet_seleccionada(self):
        r = tools_excel.leer_excel(self.xlsx, sheet_name="Totales",
                                   force_recalc=False, include_formulas=True)
        assert r.get("error") is None
        assert len(r["sheets"]) == 1
        assert r["sheets"][0]["name"] == "Totales"
        by_ref = {c["ref"]: c for c in r["sheets"][0]["cells"]}
        assert by_ref["B1"]["f"] == "=Conciliacion!D5"

    def test_max_cells_trunca(self):
        r = tools_excel.leer_excel(self.xlsx, force_recalc=False,
                                   include_formulas=False, max_cells=2)
        assert r.get("error") is None
        assert r["sheets"][0]["truncated"] is True

    def test_archivo_inexistente(self):
        r = tools_excel.leer_excel(os.path.join(self.tmp, "no_existe.xlsx"))
        assert "error" in r

    def test_formato_no_soportado(self):
        fake = os.path.join(self.tmp, "doc.csv")
        open(fake, "w").write("a,b,c\n")
        r = tools_excel.leer_excel(fake)
        assert "error" in r
        assert "csv" in r["error"].lower()

    def test_archivo_demasiado_grande(self):
        tools_excel.MAX_FILE_MB = 0.0001
        r = tools_excel.leer_excel(self.xlsx)
        assert "error" in r
        assert "excede" in r["error"].lower()

    def test_no_es_xlsx(self):
        fake = os.path.join(self.tmp, "basura.xlsx")
        open(fake, "w").write("esto no es un zip")
        r = tools_excel.leer_excel(fake)
        assert "error" in r

    def test_check_capabilities_shape(self):
        c = tools_excel.check_excel_capabilities()
        assert "platform" in c
        assert c["strategy"] in ("libreoffice", "formualizer", "cached_only")
        assert "openpyxl" in c

    def test_resumen_excel_shape_y_tipos(self):
        r = tools_excel.resumen_excel(self.xlsx, sheet_name="Conciliacion")
        assert "error" not in r
        sheet = r["sheets"][0]
        assert sheet["rows"] == 4  # filas de datos tras la cabecera (incl. la de SUM)
        assert sheet["cols"] == 7  # incluye F/G por la celda combinada "nota"
        nombres = {c["nombre"] for c in sheet["columnas"]}
        assert {"Cliente", "Base", "IVA", "Total"} <= nombres
        base_col = next(c for c in sheet["columnas"] if c["nombre"] == "Base")
        assert base_col["tipo_dominante"] == "int"
        assert base_col["nulos"] >= 1  # fila 5 (SUM) no tiene Base

    def test_resumen_excel_sugiere_agrupacion(self):
        r = tools_excel.resumen_excel(self.xlsx, sheet_name="Conciliacion")
        sugerencias = r["sheets"][0]["agrupacion_sugerida"]
        assert any("Base" in s or "IVA" in s or "Total" in s for s in sugerencias)

    def test_resumen_excel_hoja_inexistente(self):
        r = tools_excel.resumen_excel(self.xlsx, sheet_name="NoExiste")
        assert "error" in r

    def test_resumen_excel_no_modifica_original(self):
        import hashlib
        before = hashlib.sha256(open(self.xlsx, "rb").read()).hexdigest()
        tools_excel.resumen_excel(self.xlsx)
        after = hashlib.sha256(open(self.xlsx, "rb").read()).hexdigest()
        assert before == after


class TestReconciliarExcelBd(unittest.TestCase):
    def test_reconcilia_ok_y_diff(self):
        excel_rows = [
            {"codigo": "F1", "importe": 100.00},
            {"codigo": "F2", "importe": 200.00},
        ]
        bd_rows = [
            {"codigo": "F1", "importe": 100.00},
            {"codigo": "F2", "importe": 205.00},
        ]
        r = tools_excel.reconciliar_excel_bd(excel_rows, bd_rows, "codigo", "importe")
        assert r["n_ok"] == 1
        assert r["n_diff"] == 1
        diff_row = next(d for d in r["diferencias"] if d["clave"] == "F2")
        assert diff_row["estado"] == "DIFF"

    def test_reconcilia_solo_excel_y_solo_bd(self):
        excel_rows = [{"codigo": "F1", "importe": 100.00}]
        bd_rows = [{"codigo": "F2", "importe": 50.00}]
        r = tools_excel.reconciliar_excel_bd(excel_rows, bd_rows, "codigo", "importe")
        estados = {d["clave"]: d["estado"] for d in r["diferencias"]}
        assert estados == {"F1": "SOLO_EXCEL", "F2": "SOLO_BD"}

    def test_reconcilia_sin_key_col_falla(self):
        with self.assertRaises(ValueError):
            tools_excel.reconciliar_excel_bd([{"a": 1}], [{"a": 1}], "", "a")

    def test_reconcilia_fila_sin_columna_clave_falla_con_error_claro(self):
        with self.assertRaises(ValueError) as ctx:
            tools_excel.reconciliar_excel_bd([{"otra": 1}], [{"codigo": "F1"}], "codigo", "otra")
        assert "codigo" in str(ctx.exception)


if __name__ == "__main__":
    unittest.main()