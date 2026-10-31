import base64
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from mcp_bd_readonly import file_guard
from mcp_bd_readonly import tools_excel
from mcp_bd_readonly.server import excel_paths_allowed
from mcp_bd_readonly.tools_excel import leer_excel, resumen_excel


def teardown_function():
    file_guard.configure(None, allow_paths=True)


def _xlsx_b64() -> str:
    """XLSX generado en memoria con una fórmula sin valor cached."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws["A1"], ws["B1"] = "id", "importe"
    ws["A2"], ws["B2"] = 1, 10
    ws["A3"], ws["B3"] = 2, 20
    ws["B4"] = "=SUM(B2:B3)"
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


# --------------------------------------------------------------- sandbox
def test_sin_sandbox_no_restringe(tmp_path):
    file_guard.configure(None)
    assert file_guard.resolve("/etc/hosts") == Path("/etc/hosts").resolve()


def test_dentro_del_sandbox_se_permite(tmp_path):
    f = tmp_path / "a.xlsx"
    f.write_bytes(b"x")
    file_guard.configure(tmp_path)
    assert file_guard.resolve(str(f)) == f.resolve()


def test_fuera_del_sandbox_se_rechaza(tmp_path):
    file_guard.configure(tmp_path)
    with pytest.raises(file_guard.OutsideSandbox):
        file_guard.resolve("/etc/hosts")


def test_traversal_se_rechaza(tmp_path):
    root = tmp_path / "uploads"
    root.mkdir()
    (tmp_path / "secreto.xlsx").write_bytes(b"x")
    file_guard.configure(root)
    with pytest.raises(file_guard.OutsideSandbox):
        file_guard.resolve(root / ".." / "secreto.xlsx")


def test_symlink_que_escapa_se_rechaza(tmp_path):
    root = tmp_path / "uploads"
    root.mkdir()
    fuera = tmp_path / "fuera.xlsx"
    fuera.write_bytes(b"x")
    (root / "enlace.xlsx").symlink_to(fuera)
    file_guard.configure(root)
    with pytest.raises(file_guard.OutsideSandbox):
        file_guard.resolve(root / "enlace.xlsx")


def test_prefijo_parecido_no_cuela(tmp_path):
    root = tmp_path / "up"
    root.mkdir()
    gemelo = tmp_path / "uploads"
    gemelo.mkdir()
    (gemelo / "x.xlsx").write_bytes(b"x")
    file_guard.configure(root)
    with pytest.raises(file_guard.OutsideSandbox):
        file_guard.resolve(gemelo / "x.xlsx")


def test_rutas_deshabilitadas(tmp_path):
    file_guard.configure(None, allow_paths=False)
    with pytest.raises(file_guard.PathsDisabled):
        file_guard.resolve(tmp_path / "x.xlsx")


# --------------------------------------------------------------- política
def test_stdio_admite_rutas():
    assert excel_paths_allowed(False, SimpleNamespace(files_dir="")) is True


def test_remoto_sin_directorio_no_admite_rutas():
    assert excel_paths_allowed(True, SimpleNamespace(files_dir="")) is False


def test_remoto_con_directorio_admite_rutas():
    cfg = SimpleNamespace(files_dir="/var/lib/mcp-optima-bd/uploads")
    assert excel_paths_allowed(True, cfg) is True


# --------------------------------------------------------------- memoria
def test_leer_excel_base64_sin_escribir_en_disco(monkeypatch):
    file_guard.configure(None, allow_paths=False)

    def boom(**kwargs):
        raise AssertionError("no debe crearse ningún directorio temporal")

    monkeypatch.setattr(tools_excel.tempfile, "mkdtemp", boom)
    out = leer_excel(contenido_b64=_xlsx_b64())

    assert "error" not in out
    assert out["workbook"]["en_memoria"] is True
    assert out["recalc"]["engine"] == "formualizer_memory"
    assert out["recalc"]["sin_escritura"] is True
    celdas = {c["ref"]: c for c in out["sheets"][0]["cells"]}
    assert celdas["B4"]["f"] == "=SUM(B2:B3)"
    assert celdas["B4"]["v"] == 30


def test_resumen_excel_base64():
    file_guard.configure(None, allow_paths=False)
    out = resumen_excel(contenido_b64=_xlsx_b64())
    assert out["name"] == "contenido.xlsx"
    assert out["en_memoria"] is True
    assert out["sheets"][0]["name"] == "Hoja1"


def test_base64_invalido_devuelve_error():
    out = resumen_excel(contenido_b64="no es base64!!!")
    assert "error" in out


def test_sin_ruta_ni_contenido_devuelve_error():
    out = leer_excel()
    assert "error" in out
    assert "contenido_b64" in out["error"]


def test_ruta_rechazada_en_modo_servidor(tmp_path):
    file_guard.configure(None, allow_paths=False)
    f = tmp_path / "x.xlsx"
    f.write_bytes(b"x")
    out = leer_excel(str(f))
    assert "error" in out
    assert "contenido_b64" in out["error"]


def test_resumen_excel_ruta_fuera_del_sandbox(tmp_path):
    file_guard.configure(tmp_path / "uploads")
    out = resumen_excel(str(tmp_path / "otro.xlsx"))
    assert "error" in out
    assert "fuera" in out["error"]
