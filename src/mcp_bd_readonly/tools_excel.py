import base64
import binascii
import io
import json
import logging
import os
import platform
import re
import shutil
import subprocess
import tempfile
import zipfile
from openpyxl.utils import get_column_letter
from pathlib import Path

from mcp_bd_readonly import file_guard

log = logging.getLogger("mcp_bd_readonly.excel")

MAX_FILE_MB = 25
MAX_TOTAL_CELLS = 250_000
MAX_SHEET_CELLS = 60_000
MAX_CELL_TEXT = 120


class ExcelSafetyError(Exception):
    pass


def _system() -> str:
    return platform.system()


def _get_libreoffice_cmd() -> list[str]:
    if _system() == "Windows":
        candidates = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
    elif _system() == "Darwin":
        candidates = [
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
            "/opt/homebrew/bin/soffice",
            "/usr/local/bin/soffice",
        ]
    else:
        candidates = []
    for c in candidates:
        if Path(c).exists():
            return [c]
    return ["soffice"]


def _libreoffice_available() -> tuple[bool, str]:
    try:
        r = subprocess.run(
            _get_libreoffice_cmd() + ["--version"],
            capture_output=True,
            text=True,
            timeout=15,
        )
        if r.returncode == 0:
            v = (r.stdout or r.stderr).strip().splitlines()
            return True, (v[0] if v else "unknown")
    except (OSError, subprocess.TimeoutExpired):
        pass
    return False, ""


def _formualizer_available() -> tuple[bool, str]:
    try:
        import formualizer
        return True, getattr(formualizer, "__version__", "unknown")
    except Exception:
        return False, ""


def check_excel_capabilities() -> dict:
    lo_ok, lo_v = _libreoffice_available()
    fm_ok, fm_v = _formualizer_available()
    if lo_ok:
        strategy = "libreoffice"
    elif fm_ok:
        strategy = "formualizer"
    else:
        strategy = "cached_only"
    return {
        "platform": _system(),
        "openpyxl": _version("openpyxl"),
        "libreoffice": {"available": lo_ok, "version": lo_v},
        "formualizer": {"available": fm_ok, "version": fm_v},
        "strategy": strategy,
        # Ruta recomendada en servidor: bytes en memoria, sin ficheros temporales.
        "memoria": {"sin_escritura": True, "formualizer_bytes": fm_ok},
        "note": strategy if strategy != "cached_only" else
            "Sin motor de recálculo: solo valores cached. Instala LibreOffice o formualizer.",
    }


def _version(pkg: str) -> str:
    try:
        mod = __import__(pkg)
        return getattr(mod, "__version__", "unknown")
    except Exception:
        return "no disponible"


class _Source:
    """Excel ya materializado: bytes en memoria + nombre. `path` es None si vino
    en contenido_b64, y en ese caso no se escribe nada en disco en ningún paso."""

    __slots__ = ("data", "name", "path")

    def __init__(self, data: bytes, name: str, path: Path | None = None):
        self.data = data
        self.name = name
        self.path = path

    @property
    def suffix(self) -> str:
        return Path(self.name).suffix.lower()

    def stream(self) -> io.BytesIO:
        return io.BytesIO(self.data)


def _load_source(xlsx_path: str | None, contenido_b64: str | None) -> _Source:
    """Obtiene el XLSX como bytes. Nunca escribe en disco."""
    if contenido_b64:
        try:
            data = base64.b64decode(contenido_b64, validate=False)
        except (binascii.Error, ValueError) as e:
            raise ExcelSafetyError(f"contenido_b64 no es base64 válido: {e}")
        if not data:
            raise ExcelSafetyError("contenido_b64 está vacío")
        return _Source(data, "contenido.xlsx", None)
    if not xlsx_path:
        raise ExcelSafetyError(
            "indica contenido_b64 (Excel en memoria) o xlsx_path (solo en local)."
        )
    path = file_guard.resolve(xlsx_path)
    try:
        data = path.read_bytes()
    except OSError as e:
        raise ExcelSafetyError(f"No se pudo leer {path}: {e}")
    return _Source(data, path.name, path)


def _as_source(src) -> _Source:
    return src if isinstance(src, _Source) else _load_source(str(src), None)


def _safety_checks(src) -> None:
    src = _as_source(src)
    if src.suffix not in (".xlsx", ".xlsm"):
        raise ExcelSafetyError(f"Formato no soportado (solo .xlsx/.xlsm): {src.suffix}")
    size_mb = len(src.data) / (1024 * 1024)
    if size_mb > MAX_FILE_MB:
        raise ExcelSafetyError(f"Archivo excede {MAX_FILE_MB}MB ({size_mb:.1f}MB)")
    try:
        with zipfile.ZipFile(src.stream()) as z:
            bad = z.testzip()
            total = sum(i.file_size for i in z.infolist())
            if bad:
                raise ExcelSafetyError(f"XLSX corrupto (zip): {bad}")
            if total > 400 * 1024 * 1024:
                raise ExcelSafetyError(f"Descomprimir excede 400MB ({total/1e6:.0f}MB)")
    except zipfile.BadZipFile:
        raise ExcelSafetyError("No es un XLSX válido (no es ZIP)")


def _detect_formulas(path) -> dict:
    path = _as_source(path)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path.stream(), data_only=False, read_only=True)
        out = {"has_formulas": False, "per_sheet": {}, "total": 0}
        for ws in wb.worksheets:
            n = 0
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type == "f":
                        n += 1
                        out["has_formulas"] = True
                        if n >= 5000:
                            break
                if n >= 5000:
                    break
            out["per_sheet"][ws.title] = n
            out["total"] += n
        wb.close()
        return out
    except Exception as e:
        # No se asume "sin fórmulas": eso haría leer valores cached obsoletos
        # como si fueran buenos. Se marca desconocido y el llamador recalcula.
        log.warning("detect_formulas: %s", e)
        return {"has_formulas": None, "per_sheet": {}, "total": 0, "error": str(e)}


def _recalc_with_libreoffice(path: Path, tmp_dir: Path) -> dict:
    out_file = tmp_dir / f"recalc_{path.stem}.xlsx"
    flags = ["--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_dir)]
    if _system() == "Darwin":
        flags = ["--headless", "--invisible", "--norestore", "--nologo", "--calc",
                 "--convert-to", "xlsx", "--outdir", str(tmp_dir)]
    elif _system() == "Windows":
        flags = ["--headless", "--nodefault", "--norestore", "--calc",
                 "--convert-to", "xlsx", "--outdir", str(tmp_dir)]
    import subprocess as sp
    cflags = sp.CREATE_NO_WINDOW if _system() == "Windows" else 0
    try:
        r = sp.run(_get_libreoffice_cmd() + flags + [str(path)],
                   capture_output=True, text=True, timeout=120, creationflags=cflags)
        if r.returncode != 0:
            return {"ok": False, "error": (r.stderr or r.stdout or "LO returncode!=0")[:200]}
    except sp.TimeoutExpired:
        return {"ok": False, "error": "LibreOffice timeout 120s"}
    except Exception as e:
        return {"ok": False, "error": f"LibreOffice: {e}"}
    candidates = sorted(tmp_dir.glob(f"recalc_{path.stem}.xlsx")) + sorted(tmp_dir.glob("*.xlsx"))
    cand = next((c for c in candidates if c.exists() and c.stat().st_size > 0), None)
    if cand:
        return {"ok": True, "path": cand}
    return {"ok": False, "error": "LibreOffice no produjo archivo"}


def _recalc_with_formualizer(path: Path, tmp_dir: Path) -> dict:
    out_file = tmp_dir / f"recalc_{path.stem}.xlsx"
    try:
        import formualizer
        summary = formualizer.recalculate_file(str(path), output=str(out_file))
        if out_file.exists() and out_file.stat().st_size > 0:
            return {"ok": True, "path": out_file, "summary": summary}
        return {"ok": False, "error": "formualizer no produjo archivo"}
    except Exception as e:
        return {"ok": False, "error": f"formualizer: {e}"[:200]}


def _smart_recalc(path: Path, tmp_dir: Path) -> dict:
    lo_ok, _ = _libreoffice_available()
    if lo_ok:
        r = _recalc_with_libreoffice(path, tmp_dir)
        if r["ok"]:
            return {"recalculated": True, "engine": "libreoffice", "path": str(r["path"])}
        log.warning("LO recalc falló: %s", r.get("error"))
    fm_ok, _ = _formualizer_available()
    if fm_ok:
        r = _recalc_with_formualizer(path, tmp_dir)
        if r["ok"]:
            return {"recalculated": True, "engine": "formualizer", "path": str(r["path"]),
                    "warning": "formualizer: sin tablas dinámicas/LET/LAMBDA", "summary": r.get("summary")}
        log.warning("formualizer recalc falló: %s", r.get("error"))
    return {"recalculated": False, "engine": "cached_only",
            "warning": "Sin recálculo: valores cached."}


def _recalc_in_memory(src: _Source) -> dict:
    """Recalcula fórmulas sin tocar el disco (formualizer).

    Necesario cuando el XLSX llega en base64: openpyxl escribe fórmulas sin
    valor cached, así que `data_only=True` devolvería None.
    """
    try:
        import formualizer
    except Exception as e:  # noqa: BLE001 — sin motor se sigue con cached
        return {
            "recalculated": False, "engine": "cached_only", "sin_escritura": True,
            "warning": f"Sin motor de cálculo en memoria ({e}); se usan valores cached.",
        }
    try:
        book = formualizer.load_workbook_bytes(src.data)
        book.evaluate_all()
        return {
            "recalculated": True, "engine": "formualizer_memory", "sin_escritura": True,
            "book": book, "warning": "formualizer: sin tablas dinámicas/LET/LAMBDA",
        }
    except Exception as e:  # noqa: BLE001
        return {
            "recalculated": False, "engine": "cached_only", "sin_escritura": True,
            "warning": f"Recálculo en memoria falló ({e}); se usan valores cached.",
        }


def _norm(v):
    """Aplana valores para JSON, acota strings largos."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(v, 6)
    s = str(v)
    return s if len(s) <= MAX_CELL_TEXT else s[:MAX_CELL_TEXT] + "…"


_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def _sheet_xml_name(z: zipfile.ZipFile, sheet_name: str) -> str | None:
    """Resuelve el XML de una hoja concreta: workbook.xml → r:id → rels."""
    import xml.etree.ElementTree as ET

    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rid = None
    for sh in wb.iter(f"{{{_NS_MAIN}}}sheet"):
        if sh.get("name") == sheet_name:
            rid = sh.get(f"{{{_NS_DOC_REL}}}id")
            break
    if not rid:
        return None
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    for rel in rels.iter(f"{{{_NS_PKG_REL}}}Relationship"):
        if rel.get("Id") == rid:
            target = (rel.get("Target") or "").lstrip("/")
            if not target:
                return None
            return target if target.startswith("xl/") else f"xl/{target}"
    return None


def _merged_ranges(path, sheet_name: str) -> list[str]:
    """Merged ranges de ESA hoja (read_only no los expone).

    Se resuelve nombre → XML de hoja. Antes se devolvía el primer bloque
    `mergeCells` del workbook, así que las hojas 2+ recibían los rangos de la 1.
    """
    path = _as_source(path)
    try:
        import xml.etree.ElementTree as ET

        with zipfile.ZipFile(path.stream()) as z:
            name = _sheet_xml_name(z, sheet_name)
            if not name:
                return []
            root = ET.fromstring(z.read(name))
            sh = root.find(f".//{{{_NS_MAIN}}}mergeCells")
            if sh is None:
                return []
            return [mc.attrib["ref"] for mc in sh.findall(f"{{{_NS_MAIN}}}mergeCell")]
    except Exception as e:  # noqa: BLE001 — un XLSX raro no debe tumbar la lectura
        log.warning("merged_ranges(%s): %s", sheet_name, e)
        return []


def _mem_value(book, sheet: str, row: int, col: int):
    """Valor calculado en memoria por formualizer para una celda con fórmula."""
    if book is None:
        return None
    try:
        return book.get_value(sheet, row, col)
    except Exception as e:  # noqa: BLE001 — una celda rara no debe tumbar la hoja
        log.debug("mem get_value %s (%s,%s): %s", sheet, row, col, e)
        return None


def _sheet_to_json(ws, formulas: dict[str, str], max_cells: int, merged: list[str],
                   mem=None) -> dict:
    cells = []
    for r_i, row in enumerate(ws.iter_rows(), start=1):
        for c_i, c in enumerate(row, start=1):
            if not hasattr(c, "value"):
                continue
            v = c.value
            coord = get_column_letter(c_i) + str(r_i)
            f = formulas.get(coord)
            if v is None and f and mem is not None:
                # Sin valor cached: se usa el cálculo hecho en memoria.
                v = _mem_value(mem, ws.title, r_i, c_i)
            if v is None and not f:
                continue
            cell = {"ref": coord}
            if f:
                cell["f"] = f[:200]
            if v is not None:
                cell["v"] = _norm(v)
            cells.append(cell)
            if len(cells) >= max_cells:
                return {
                    "name": ws.title,
                    "rows": ws.max_row, "cols": ws.max_column,
                    "cells": cells,
                    "merged_ranges": merged,
                    "truncated": True,
                }
    return {
        "name": ws.title,
        "rows": ws.max_row, "cols": ws.max_column,
        "cells": cells,
        "merged_ranges": merged,
        "truncated": False,
    }


def _formula_map(wb, sheet_name: str | None) -> dict[str, dict[str, str]]:
    from openpyxl import load_workbook
    out = {}
    wbf = load_workbook(wb, data_only=False, read_only=True)
    for ws in wbf.worksheets:
        if sheet_name and ws.title != sheet_name:
            continue
        m = {}
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f" and c.value and str(c.value).startswith("="):
                    m[c.coordinate] = str(c.value)
        out[ws.title] = m
        if sheet_name:
            break
    wbf.close()
    return out


def _resumen_sheet(ws, max_rows_preview: int, max_rows_scan: int) -> dict:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return {"name": ws.title, "rows": 0, "cols": 0, "columnas": [], "preview": [],
                "agrupacion_sugerida": []}
    headers = [str(h) if h is not None else f"col_{i + 1}" for i, h in enumerate(header)]
    stats = [{"nulos": 0, "tipos": {}} for _ in headers]
    preview = []
    n_rows = 0
    for row in rows_iter:
        n_rows += 1
        if n_rows > max_rows_scan:
            break
        if len(preview) < max_rows_preview:
            preview.append([_norm(v) for v in row[:len(headers)]])
        for i, v in enumerate(row[:len(headers)]):
            if v is None:
                stats[i]["nulos"] += 1
            else:
                tname = type(v).__name__
                stats[i]["tipos"][tname] = stats[i]["tipos"].get(tname, 0) + 1

    columnas = []
    fechas, numericas = [], []
    for h, s in zip(headers, stats):
        tipo_dominante = max(s["tipos"], key=s["tipos"].get) if s["tipos"] else "vacio"
        columnas.append({"nombre": h, "nulos": s["nulos"], "tipos": s["tipos"],
                          "tipo_dominante": tipo_dominante})
        hl = h.lower()
        if "fecha" in hl or tipo_dominante in ("date", "datetime"):
            fechas.append(h)
        elif tipo_dominante in ("int", "float", "Decimal") or any(
                k in hl for k in ("importe", "total", "base", "iva", "monto")):
            numericas.append(h)

    agrupacion = []
    if fechas and numericas:
        agrupacion.append(f"Agrupar por {fechas[0]} (mes/año) sumando {', '.join(numericas)}")
    elif numericas:
        agrupacion.append(f"Sumar/agregar {', '.join(numericas)}")

    return {"name": ws.title, "rows": n_rows, "cols": len(headers),
            "columnas": columnas, "preview": preview, "agrupacion_sugerida": agrupacion}


def resumen_excel(xlsx_path: str | None = None, sheet_name: str | None = None,
                   max_rows_preview: int = 10, max_rows_scan: int = 5000,
                   contenido_b64: str | None = None) -> dict:
    """Shape de un .xlsx sin exigir recálculo: dims, tipos por columna, nulos,
    primeras filas y una sugerencia simple de agrupación. Usa valores cached (data_only),
    nunca modifica el original.

    `contenido_b64`: Excel en base64, procesado **en memoria** (recomendado en
    servidor: no se escribe nada en disco).
    """
    try:
        src = _load_source(xlsx_path, contenido_b64)
        _safety_checks(src)
    except (file_guard.OutsideSandbox, file_guard.PathsDisabled, ExcelSafetyError) as e:
        return {"error": str(e)}

    from openpyxl import load_workbook
    wb = load_workbook(src.stream(), data_only=True, read_only=True)
    try:
        out_sheets = []
        for ws in wb.worksheets:
            if sheet_name and ws.title != sheet_name:
                continue
            out_sheets.append(_resumen_sheet(ws, max_rows_preview, max_rows_scan))
            if sheet_name:
                break
    finally:
        wb.close()

    if sheet_name and not out_sheets:
        return {"error": f"Hoja no encontrada: {sheet_name}"}
    return {"name": src.name, "en_memoria": src.path is None, "sheets": out_sheets}


def reconciliar_excel_bd(excel_rows: list[dict], bd_rows: list[dict], key_col: str,
                          importe_col: str, fecha_col: str | None = None,
                          tolerancia: float = 0.01) -> dict:
    """Concilia filas de una hoja Excel (ya leída, p.ej. con leer_excel/resumen_excel)
    contra filas de la BD (ya obtenidas, p.ej. con run_query), cruzando por key_col.
    key_col es OBLIGATORIO y explícito: no intenta adivinar una clave difusa.
    Devuelve diff por línea (OK/DIFF/SOLO_EXCEL/SOLO_BD) + totales."""
    if not key_col:
        raise ValueError("key_col es obligatorio: indica qué columna identifica cada línea de forma única")

    from decimal import Decimal

    def _idx(rows, origen):
        # Clave → lista de filas. Antes se sobreescribía: dos líneas con el mismo
        # código colapsaban en una y el descuadre desaparecía del informe.
        out: dict = {}
        for r in rows:
            if key_col not in r:
                raise ValueError(f"Falta la columna clave '{key_col}' en una fila de {origen}")
            out.setdefault(r[key_col], []).append(r)
        return out

    def _total(rows) -> Decimal:
        total = Decimal("0")
        for r in rows:
            total += Decimal(str(r[importe_col] or 0))
        return total

    excel_idx = _idx(excel_rows, "excel")
    bd_idx = _idx(bd_rows, "bd")
    tol = Decimal(str(tolerancia))

    diferencias = []
    total_excel = Decimal("0")
    total_bd = Decimal("0")
    n_duplicados = 0
    for key in sorted(set(excel_idx) | set(bd_idx), key=str):
        e_rows = excel_idx.get(key)
        b_rows = bd_idx.get(key)
        importe_e = _total(e_rows) if e_rows else None
        importe_b = _total(b_rows) if b_rows else None
        if importe_e is not None:
            total_excel += importe_e
        if importe_b is not None:
            total_bd += importe_b

        if b_rows is None:
            diferencias.append({"clave": key, "estado": "SOLO_EXCEL",
                                 "importe_excel": importe_e,
                                 "n_lineas_excel": len(e_rows)})
            continue
        if e_rows is None:
            diferencias.append({"clave": key, "estado": "SOLO_BD",
                                 "importe_bd": importe_b,
                                 "n_lineas_bd": len(b_rows)})
            continue

        diff = importe_e - importe_b
        item = {"clave": key, "importe_excel": importe_e, "importe_bd": importe_b,
                "n_lineas_excel": len(e_rows), "n_lineas_bd": len(b_rows),
                "diff": diff}
        if fecha_col:
            item["fecha_excel"] = e_rows[0].get(fecha_col)
            item["fecha_bd"] = b_rows[0].get(fecha_col)
        if len(e_rows) > 1 or len(b_rows) > 1:
            item["estado"] = "DUPLICADO"
            n_duplicados += 1
        else:
            item["estado"] = "OK" if abs(diff) <= tol else "DIFF"
        diferencias.append(item)

    return {
        "diferencias": diferencias,
        "totales": {"excel": total_excel, "bd": total_bd, "diff": total_excel - total_bd},
        "n_ok": sum(1 for d in diferencias if d["estado"] == "OK"),
        "n_diff": sum(
            1 for d in diferencias
            if d["estado"] in ("DIFF", "SOLO_EXCEL", "SOLO_BD", "DUPLICADO")
        ),
        "n_duplicados": n_duplicados,
    }


def leer_excel(xlsx_path: str | None = None, sheet_name: str | None = None,
               include_formulas: bool = True, force_recalc: bool = True,
               max_cells: int = 5000, contenido_b64: str | None = None) -> dict:
    """Lee un XLSX de finanzas y devuelve JSON consumible por un LLM para auditoría.

    Nunca modifica el original. `contenido_b64` procesa el Excel **en memoria**
    (nada se escribe en disco: el recálculo lo hace formualizer sobre los bytes).
    Con `xlsx_path` (solo local) se mantiene el recálculo por LibreOffice en un
    temporal que se borra al terminar. max_cells limita celdas por hoja.
    """
    try:
        src = _load_source(xlsx_path, contenido_b64)
        _safety_checks(src)
    except (file_guard.OutsideSandbox, file_guard.PathsDisabled, ExcelSafetyError) as e:
        return {"error": str(e)}

    from openpyxl import load_workbook

    formula_info = _detect_formulas(src)
    formulas_unknown = formula_info.get("has_formulas") is None
    recalc = {"recalculated": False, "engine": "cached_only"}
    mem = None
    tmp_dir = None
    read_stream = src.stream()

    if force_recalc and (formula_info.get("has_formulas") or formulas_unknown):
        if src.path is None:
            # En memoria no se escribe nada: recálculo sobre los bytes.
            rec = _recalc_in_memory(src)
            mem = rec.pop("book", None)
            recalc = rec
        else:
            tmp_dir = Path(tempfile.mkdtemp(prefix="xlsxrecalc_"))
            recalc = _smart_recalc(src.path, tmp_dir)
            if recalc.get("recalculated"):
                read_stream = io.BytesIO(Path(recalc["path"]).read_bytes())

    try:
        wb = load_workbook(read_stream, data_only=True, read_only=True)
        fmaps = {}
        if include_formulas and formula_info.get("has_formulas"):
            fmaps = _formula_map(src.stream(), sheet_name)
        sheets = []
        for ws in wb.worksheets:
            if sheet_name and ws.title != sheet_name:
                continue
            fmap = fmaps.get(ws.title, {})
            sheets.append(_sheet_to_json(ws, fmap, min(max_cells, MAX_SHEET_CELLS),
                                         _merged_ranges(src, ws.title), mem))
            if sheet_name:
                break
        meta = {
            "name": src.name,
            "en_memoria": src.path is None,
            "platform": _system(),
            "sheets_meta": [
                {"name": ws.title, "rows": ws.max_row, "cols": ws.max_column,
                 "formulas": formula_info["per_sheet"].get(ws.title, 0)}
                for ws in wb.worksheets
            ],
        }
        wb.close()
    except Exception as e:
        if tmp_dir:
            shutil.rmtree(tmp_dir, ignore_errors=True)
        return {"error": f"No se pudo leer XLSX: {e}"}

    if tmp_dir:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if sheet_name:
        chosen = next((s for s in sheets if s["name"] == sheet_name), None)
        sheets = [chosen] if chosen else []

    out = {
        "workbook": meta,
        "recalc": recalc,
        "has_formulas": formula_info.get("has_formulas", False),
        "formula_count": formula_info.get("total", 0),
        "sheets": sheets,
    }
    if formula_info.get("error"):
        out["formulas_error"] = formula_info["error"]
    return out